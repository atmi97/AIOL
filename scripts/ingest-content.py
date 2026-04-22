#!/usr/bin/env python3
"""Convert the three source docs into structured content for the platform.

Inputs:  source-docs/3sHealth-AIOL-Program-Plan-v2.1.docx
         source-docs/AIOL-Tier1-Content-Package-v1.docx
         source-docs/AIOL-Tier1-Quiz-Bank-v1.xlsx

Outputs: content/program.mdx           (from program plan)
         content/tier1/meta.json
         content/tier1/module-1.1/index.mdx
         content/tier1/module-1.2/index.mdx
         content/tier1/module-1.3/index.mdx
         content/tier1/module-1.4/index.mdx
         content/tier1/module-1.5/index.mdx
         content/tier1/quiz-bank.json

Run: python3 scripts/ingest-content.py
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "source-docs"
CONTENT = ROOT / "content"
TIER1 = CONTENT / "tier1"


def iter_block_items(parent) -> Iterable:
    """Yield paragraphs and tables in document order."""
    from docx.oxml.ns import qn

    for child in parent.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def slugify(s: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s.strip().lower())
    return s.strip("-")


def style_level(style_name: str) -> int | None:
    m = re.match(r"Heading (\d)", style_name or "")
    return int(m.group(1)) if m else None


def table_to_markdown(tbl: Table) -> str:
    rows = [[c.text.strip().replace("\n", "<br />") for c in row.cells] for row in tbl.rows]
    if not rows:
        return ""
    widths = [max(len(r[i]) for r in rows) for i in range(len(rows[0]))]
    # If first row looks like a header (no explicit style info available here),
    # render as markdown table. Otherwise keep as plain bold-labelled block.
    if all(rows[0]) and len(rows) > 1:
        head = "| " + " | ".join(rows[0]) + " |"
        sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
        body = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
        return f"{head}\n{sep}\n{body}\n"
    # Fallback: key/value block
    lines = []
    for r in rows:
        lines.append("**" + r[0] + "** — " + " ".join(r[1:]) if len(r) > 1 else "**" + r[0] + "**")
    return "\n\n".join(lines) + "\n"


@dataclass
class Module:
    id: str                # "1.1"
    slug: str              # "module-1.1"
    title: str             # "What Is AI?"
    duration_min: int
    objectives: list[str]
    body_mdx: str


def parse_tier1_content() -> tuple[list[Module], dict]:
    """Parse AIOL-Tier1-Content-Package-v1.docx into modules."""
    doc = Document(SRC / "AIOL-Tier1-Content-Package-v1.docx")
    modules: list[Module] = []
    appendix_blocks: list[str] = []

    current_section: str | None = None  # "module" | "appendix" | None
    module_id: str | None = None
    module_title: str | None = None
    duration_min: int = 0
    objectives: list[str] = []
    in_objectives = False
    body_lines: list[str] = []

    def flush_module() -> None:
        nonlocal module_id, module_title, duration_min, objectives, body_lines
        if module_id and module_title:
            slug = f"module-{module_id}"
            body = "\n".join(body_lines).strip() + "\n"
            modules.append(
                Module(
                    id=module_id,
                    slug=slug,
                    title=module_title,
                    duration_min=duration_min,
                    objectives=list(objectives),
                    body_mdx=body,
                )
            )
        module_id = None
        module_title = None
        duration_min = 0
        objectives = []
        body_lines.clear()

    for block in iter_block_items(doc):
        if isinstance(block, Table):
            md = table_to_markdown(block)
            if current_section == "module":
                body_lines.append(md)
            elif current_section == "appendix":
                appendix_blocks.append(md)
            continue

        text = block.text.strip()
        style = block.style.name if block.style else ""
        level = style_level(style)

        if level == 1:
            m = re.match(r"Module\s+(\d+\.\d+)\s*[—-]\s*(.+)", text)
            if m:
                flush_module()
                current_section = "module"
                module_id = m.group(1)
                module_title = m.group(2).strip()
                in_objectives = False
                continue
            if text.startswith("Appendix"):
                flush_module()
                current_section = "appendix"
                appendix_blocks.append(f"## {text}\n")
                continue
            # Other H1s
            if current_section == "module":
                body_lines.append(f"## {text}\n")
            elif current_section == "appendix":
                appendix_blocks.append(f"## {text}\n")
            continue

        if current_section == "module":
            # Duration line appears right after module heading.
            if text.lower().startswith("duration:"):
                m = re.search(r"(\d+)\s*minute", text)
                if m:
                    duration_min = int(m.group(1))
                continue

            # Learning objectives section
            if level == 3 and text.lower() == "learning objectives":
                in_objectives = True
                continue
            if level in (2, 3) and in_objectives and text.lower() != "learning objectives":
                in_objectives = False

            if in_objectives and style == "List Paragraph" and text:
                objectives.append(text.rstrip("."))
                continue
            if in_objectives and text.lower().startswith("by the end of this module"):
                continue

            # Body content
            if level == 2:
                body_lines.append(f"### {text}\n")
            elif level == 3:
                body_lines.append(f"#### {text}\n")
            elif style == "List Paragraph" and text:
                body_lines.append(f"- {text}")
            elif text:
                body_lines.append(text + "\n")
            else:
                body_lines.append("")
        elif current_section == "appendix":
            if level == 2:
                appendix_blocks.append(f"### {text}\n")
            elif level == 3:
                appendix_blocks.append(f"#### {text}\n")
            elif style == "List Paragraph" and text:
                appendix_blocks.append(f"- {text}")
            elif text:
                appendix_blocks.append(text + "\n")

    flush_module()

    tier_meta = {
        "id": "tier1",
        "title": "Tier 1 — AI Awareness",
        "tagline": "Foundational certification in the AMS AI Operator Licence Program",
        "audience": "All AMS staff",
        "estMinutes": 120,
        "passThreshold": 0.8,
        "renewal": "Annual",
        "philosophy": "Human verify, Human decide, Human accountable.",
        "modules": [
            {"id": m.id, "slug": m.slug, "title": m.title, "estMinutes": m.duration_min, "order": i + 1}
            for i, m in enumerate(sorted(modules, key=lambda x: x.id))
        ],
        "appendix": "\n".join(appendix_blocks),
    }

    return modules, tier_meta


def parse_program_plan() -> str:
    """Render the program plan as a single MDX doc (read-only reference)."""
    doc = Document(SRC / "3sHealth-AIOL-Program-Plan-v2.1.docx")
    lines: list[str] = ["---", "title: AMS AI Operator Licence Program", "type: reference", "---", ""]
    for block in iter_block_items(doc):
        if isinstance(block, Table):
            lines.append(table_to_markdown(block))
            continue
        text = block.text.rstrip()
        if not text.strip():
            lines.append("")
            continue
        style = block.style.name if block.style else ""
        level = style_level(style)
        if level:
            lines.append("#" * min(level, 4) + " " + text)
        elif style == "List Paragraph":
            lines.append("- " + text)
        else:
            lines.append(text)
    return "\n".join(lines).strip() + "\n"


LETTER_TO_ID = {"A": "a", "B": "b", "C": "c", "D": "d"}


def parse_quiz_bank() -> dict:
    """Parse AIOL-Tier1-Quiz-Bank-v1.xlsx into structured JSON."""
    wb = load_workbook(SRC / "AIOL-Tier1-Quiz-Bank-v1.xlsx", data_only=True)

    ws = wb["Question Bank"]
    headers = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col["ID"]]:
            continue
        correct_letter = (row[col["Correct"]] or "").strip().upper()
        options = []
        for letter in ["A", "B", "C", "D"]:
            text = (row[col[letter]] or "").strip()
            if text:
                options.append({"id": LETTER_TO_ID[letter], "text": text})
        q = {
            "id": row[col["ID"]],
            "module": str(row[col["Module"]]).strip(),
            "moduleRef": f"module-{str(row[col['Module']]).strip()}",
            "lo": (row[col["LO"]] or "").strip(),
            "difficulty": (row[col["Difficulty"]] or "Medium").strip(),
            "tags": [t.strip() for t in (row[col["Tags"]] or "").split(",") if t.strip()],
            "type": "single",
            "prompt": (row[col["Question"]] or "").strip(),
            "options": options,
            "correct": [LETTER_TO_ID[correct_letter]] if correct_letter in LETTER_TO_ID else [],
            "rationale": (row[col["Rationale"]] or "").strip(),
        }
        questions.append(q)

    # Per-module blueprint: how many questions per attempt.
    module_distribution = {
        "1.1": 4,
        "1.2": 3,
        "1.3": 5,
        "1.4": 5,
        "1.5": 3,
    }

    quiz = {
        "id": "t1-final",
        "tierId": "tier1",
        "title": "Tier 1 Final Assessment",
        "description": (
            "20 randomized multiple-choice questions drawn from a bank of 60. "
            "80% pass (16 of 20). 3 attempts. 30-minute time limit."
        ),
        "passThreshold": 0.8,
        "maxAttempts": 3,
        "timeLimitMinutes": 30,
        "questionsPerAttempt": 20,
        "shuffle": True,
        "shuffleOptions": True,
        "moduleDistribution": module_distribution,
        "questions": questions,
    }

    return {"quizzes": [quiz]}


def write_module_mdx(path: Path, m: Module, tier_id: str, order: int, total: int) -> None:
    frontmatter = [
        "---",
        f"id: {m.slug}",
        f"moduleNumber: \"{m.id}\"",
        f"title: \"{m.title}\"",
        f"tierId: {tier_id}",
        f"order: {order}",
        f"estMinutes: {m.duration_min}",
        "objectives:",
    ]
    for o in m.objectives:
        safe = o.replace('"', '\\"')
        frontmatter.append(f"  - \"{safe}\"")
    frontmatter.append("---")
    frontmatter.append("")
    frontmatter.append(f"# Module {m.id} — {m.title}")
    frontmatter.append("")
    path.write_text("\n".join(frontmatter) + "\n" + m.body_mdx, encoding="utf-8")


def main() -> None:
    CONTENT.mkdir(parents=True, exist_ok=True)
    TIER1.mkdir(parents=True, exist_ok=True)

    modules, tier_meta = parse_tier1_content()
    modules.sort(key=lambda x: x.id)

    # Per-module directories
    for i, m in enumerate(modules, start=1):
        mod_dir = TIER1 / m.slug
        mod_dir.mkdir(exist_ok=True)
        write_module_mdx(mod_dir / "index.mdx", m, "tier1", i, len(modules))

    # Tier meta
    (TIER1 / "meta.json").write_text(json.dumps(tier_meta, indent=2, ensure_ascii=False), encoding="utf-8")

    # Quiz bank
    quiz_data = parse_quiz_bank()
    (TIER1 / "quiz-bank.json").write_text(json.dumps(quiz_data, indent=2, ensure_ascii=False), encoding="utf-8")

    # Program plan as reference doc
    (CONTENT / "program.mdx").write_text(parse_program_plan(), encoding="utf-8")

    print(f"Wrote {len(modules)} modules, meta.json, quiz-bank.json, program.mdx")
    for m in modules:
        print(f"  - {m.slug}: {m.title} ({m.duration_min} min, {len(m.objectives)} objectives)")
    print(f"Quiz bank: {len(quiz_data['quizzes'][0]['questions'])} questions")


if __name__ == "__main__":
    main()
